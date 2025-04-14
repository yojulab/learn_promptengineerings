from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import os

class PPTCreator:
    def __init__(self):
        self.prs = Presentation()
        self.output_dir = "codes/create_PPTs/output"
        
    def create_title_slide(self, title, subtitle=None):
        """제목 슬라이드 생성"""
        layout = self.prs.slide_layouts[0]
        slide = self.prs.slides.add_slide(layout)
        
        title_placeholder = slide.shapes.title
        subtitle_placeholder = slide.placeholders[1]
        
        title_placeholder.text = title
        if subtitle:
            subtitle_placeholder.text = subtitle
            
    def create_content_slide(self, title, content):
        """내용 슬라이드 생성"""
        layout = self.prs.slide_layouts[1]
        slide = self.prs.slides.add_slide(layout)
        
        title_placeholder = slide.shapes.title
        content_placeholder = slide.placeholders[1]
        
        title_placeholder.text = title
        content_placeholder.text = content
        
    def add_image_slide(self, title, image_path):
        """이미지 슬라이드 생성"""
        layout = self.prs.slide_layouts[5]
        slide = self.prs.slides.add_slide(layout)
        
        title_placeholder = slide.shapes.title
        title_placeholder.text = title
        
        # 이미지 추가
        left = Inches(2)
        top = Inches(2)
        if os.path.exists(image_path):
            slide.shapes.add_picture(image_path, left, top)
            
    def save(self, filename):
        """프레젠테이션 저장"""
        if not os.path.exists(self.output_dir):
            os.makedirs(self.output_dir)
        
        output_path = os.path.join(self.output_dir, filename)
        self.prs.save(output_path)
        print(f"프레젠테이션이 저장되었습니다: {output_path}")

def main():
    # 예제 사용법
    creator = PPTCreator()
    
    # 제목 슬라이드 생성
    creator.create_title_slide(
        "발표 제목",
        "부제목"
    )
    
    # 내용 슬라이드 생성
    creator.create_content_slide(
        "섹션 제목",
        "• 첫 번째 항목\n• 두 번째 항목"
    )
    
    # 이미지 슬라이드 생성료
    creator.add_image_slide(
        "이미지 제목",
        "이미지경로.jpg"
    )
    
    # 저장
    creator.save("발표자료.pptx")

if __name__ == "__main__":
    main() 