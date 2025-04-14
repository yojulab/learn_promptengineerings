from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def apply_text_formatting(shape, font_size=18, font_name="맑은 고딕", 
                        color=RGBColor(0, 0, 0), bold=False, 
                        alignment=PP_ALIGN.LEFT):
    """텍스트 서식 적용"""
    text_frame = shape.text_frame
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = alignment
    
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    font = run.font
    font.size = Pt(font_size)
    font.name = font_name
    font.color.rgb = color
    font.bold = bold

def set_slide_background(slide, rgb_color):
    """슬라이드 배경색 설정"""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*rgb_color)

def adjust_image_size(slide, picture, max_width=Inches(8), max_height=Inches(5)):
    """이미지 크기 자동 조정"""
    aspect_ratio = picture.height / picture.width
    
    if picture.width > max_width:
        picture.width = max_width
        picture.height = max_width * aspect_ratio
        
    if picture.height > max_height:
        picture.height = max_height
        picture.width = max_height / aspect_ratio

def create_bullet_list(text_frame, items, level=0):
    """글머리 기호 목록 생성"""
    for item in items:
        p = text_frame.add_paragraph()
        p.text = item
        p.level = level 